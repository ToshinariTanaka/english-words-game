#!/usr/bin/env python3
"""Prepare Nova/Ash jobs offline. This program NEVER calls an API or uploads data.

The plan is deliberately not a playback manifest: only a later generation,
validation and publication step may make an asset available to students.
"""
import argparse
import csv
import hashlib
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

MODES = {
    "word": ("★英単語", "w"), "chunk": ("★チャンク", "c"),
    "phrase": ("★文節和訳", "p"), "definition": ("★英文和訳", "s"),
}
MODEL = "gpt-4o-mini-tts-2025-12-15"
VOICES = ("nova", "ash")
PROFILE = "up-us-neutral-v1"
PROCESSING = "mono-mp3-96k-review-v1"
INSTRUCTIONS = (
    "Read only the supplied English text, exactly once, in clear General American English. "
    "Use a natural, unhurried pace and a neutral educational tone with minimal acting. "
    "Do not add introductions, explanations, translations or sound effects. "
    "For an isolated word, pronounce the word naturally; do not spell out its letters."
)
REVIEW_WORDS = {"a", "the", "live", "minute", "read", "record", "present", "use", "wind", "lead", "bow", "close"}


def canonical_text(value):
    return unicodedata.normalize("NFC", str(value or "").strip())


def digest(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def collect_items(path):
    """A single streaming pass, with errors instead of silently skipped questions."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    items, seen = [], set()
    try:
        for mode, (sheet, prefix) in MODES.items():
            if sheet not in workbook.sheetnames:
                raise ValueError("必須シートがありません: " + sheet)
            rows = workbook[sheet].iter_rows(values_only=True)
            header = next(rows, ())
            if len(header) < 13 or header[2] != "question" or header[12] != "question_key":
                raise ValueError("C列question／M列question_keyを確認してください: " + sheet)
            for row_number, row in enumerate(rows, 2):
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                if len(row) < 13 or not all(canonical_text(row[i]) for i in (2, 3, 4, 5, 6, 12)):
                    raise ValueError(f"必須欄が空です: {sheet} 行{row_number}")
                key, text = canonical_text(row[12]), canonical_text(row[2])
                if not re.fullmatch(prefix + r"\d{6}", key) or key in seen:
                    raise ValueError(f"不正または重複する問題キー: {sheet} 行{row_number} {key}")
                if canonical_text(row[1]) not in ("A1", "A2", "B1", "B2", "C1", "C2"):
                    raise ValueError(f"レベルが不正です: {key}")
                if len(text) > 4096:
                    raise ValueError(f"読み上げ入力が長すぎます: {key}")
                seen.add(key)
                items.append({"mode": mode, "question_key": key, "text": text,
                              "text_sha256": digest(text), "excel_row": row_number,
                              "pronunciation_review": mode == "word" and text.lower() in REVIEW_WORDS})
    finally:
        workbook.close()
    return items


def sample_items(items):
    """30 questions: 12 words and 6 from each other mode, including long texts."""
    chosen = []
    for mode in MODES:
        pool = sorted((i for i in items if i["mode"] == mode), key=lambda i: (len(i["text"]), i["question_key"]))
        limit = min(12 if mode == "word" else 6, len(pool))
        selected = []
        if not limit:
            continue
        # Include shortest and longest before pronunciation cases.
        for item in [pool[0], pool[-1]] + [i for i in pool if i["pronunciation_review"]]:
            if item not in selected and len(selected) < limit:
                selected.append(item)
        for n in range(limit):
            item = pool[round(n * (len(pool) - 1) / max(1, limit - 1))]
            if item not in selected and len(selected) < limit:
                selected.append(item)
        for item in pool:
            if item not in selected and len(selected) < limit:
                selected.append(item)
        chosen.extend(sorted(selected, key=lambda i: i["question_key"]))
    return chosen


def make_plan(items, source_sha256, overrides=None):
    overrides = overrides or {}
    if not isinstance(overrides, dict) or any(not isinstance(v, str) or not v.strip() for v in overrides.values()):
        raise ValueError("発音指定は問題キーと空でない英文指示のJSONオブジェクトにしてください。")
    unknown = set(overrides) - {i["question_key"] for i in items}
    if unknown:
        raise ValueError("存在しない発音指定キー: " + ", ".join(sorted(unknown)))
    jobs = []
    for item in items:
        pronunciation = overrides.get(item["question_key"], "")
        instructions = INSTRUCTIONS + (" Pronunciation guidance: " + pronunciation if pronunciation else "")
        for voice in VOICES:
            recipe = {"mode": item["mode"], "question_key": item["question_key"],
                      "text_sha256": item["text_sha256"], "voice": voice, "model": MODEL,
                      "instructions": instructions, "profile": PROFILE, "processing": PROCESSING}
            asset_id = digest(json.dumps(recipe, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
            jobs.append({**item, **recipe, "asset_id": asset_id, "status": "planned",
                         "object_key": f"audio/junior/{voice}/{item['mode']}/{item['question_key']}/{asset_id}.mp3",
                         "pronunciation_review": item["pronunciation_review"] and not bool(pronunciation)})
    return {"schema_version": 1, "kind": "generation_plan", "source_sha256": source_sha256,
            "model_shutdown_date": "2027-01-06", "requires_model_and_budget_check": True,
            "question_count": len(items), "job_count": len(jobs),
            "input_characters_per_voice": sum(len(i["text"]) for i in items), "jobs": jobs}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--pronunciations", type=Path, help="問題キーごとの英語の発音指定JSON")
    args = parser.parse_args()
    try:
        items = collect_items(args.workbook)
        overrides = json.loads(args.pronunciations.read_text(encoding="utf-8")) if args.pronunciations else {}
        plan = make_plan(items, hashlib.sha256(args.workbook.read_bytes()).hexdigest(), overrides)
        sample_keys = {i["question_key"] for i in sample_items(items)}
        sample = {**plan, "jobs": [j for j in plan["jobs"] if j["question_key"] in sample_keys]}
        sample["question_count"] = len(sample_keys)
        sample["job_count"] = len(sample["jobs"])
        sample["input_characters_per_voice"] = sum(len(i["text"]) for i in items if i["question_key"] in sample_keys)
        args.output.mkdir(parents=True, exist_ok=True)
        for filename, data in (("generation-plan.json", plan), ("sample-plan.json", sample)):
            temporary = args.output / (filename + ".tmp")
            temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            temporary.replace(args.output / filename)
        with (args.output / "sample-review.csv").open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["mode", "question_key", "text", "voice", "pronunciation_review", "status"])
            writer.writeheader()
            writer.writerows({k: job[k] for k in writer.fieldnames} for job in sample["jobs"])
        print(json.dumps({"questions": len(items), "jobs": len(plan["jobs"]),
                          "sample_jobs": len(sample["jobs"]), "paid_api_calls": 0}, ensure_ascii=False))
    except (ValueError, OSError) as error:
        parser.error(str(error))


if __name__ == "__main__":
    main()
