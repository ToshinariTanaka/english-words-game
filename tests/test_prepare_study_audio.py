import copy
import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

spec = importlib.util.spec_from_file_location("prepare_audio", Path(__file__).parents[1] / "tools/prepare_study_audio.py")
audio = importlib.util.module_from_spec(spec)
spec.loader.exec_module(audio)


class PreparationTests(unittest.TestCase):
    def workbook(self, path):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for mode, (sheet, prefix) in audio.MODES.items():
            ws = wb.create_sheet(sheet)
            ws.append(["row_number", "level", "question", "correct", "choice1", "choice2", "choice3", "", "", "", "", "", "question_key"])
            for n in range(1, 21):
                ws.append([n, "A1", "read" if n == 1 else "word " * n, "正解", "誤答1", "誤答2", "誤答3", None, None, None, None, None, f"{prefix}{n:06d}"])
        wb.save(path)

    def test_offline_two_voices_stable_ids_and_changes(self):
        with tempfile.TemporaryDirectory() as directory, patch("urllib.request.urlopen", side_effect=AssertionError("no network")):
            path = Path(directory) / "words.xlsx"
            self.workbook(path)
            items = audio.collect_items(path)
            plan = audio.make_plan(items, "source")
            self.assertEqual(plan["job_count"], 160)
            self.assertEqual(len({j["object_key"] for j in plan["jobs"]}), 160)
            self.assertEqual({j["voice"] for j in plan["jobs"]}, {"nova", "ash"})
            self.assertTrue(all(j["status"] == "planned" for j in plan["jobs"]))
            # Row order and Japanese choices do not alter audio identities.
            changed = copy.deepcopy(items)
            changed[0]["excel_row"] = 999
            self.assertEqual(audio.make_plan(changed, "other")["jobs"][0]["asset_id"], plan["jobs"][0]["asset_id"])
            override = audio.make_plan(items, "source", {"w000001": "Use present-tense /riːd/."})
            self.assertNotEqual(override["jobs"][0]["asset_id"], plan["jobs"][0]["asset_id"])
            self.assertFalse(override["jobs"][0]["pronunciation_review"])
            self.assertEqual(len(audio.sample_items(items)), 30)
            self.assertEqual(len({i["question_key"] for i in audio.sample_items(items)}), 30)

    def test_invalid_and_duplicate_keys_are_reported_not_silently_skipped(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "words.xlsx"
            for value in ["w000001", "c000002", None]:
                self.workbook(path)
                wb = openpyxl.load_workbook(path)
                wb["★英単語"]["M3"] = value
                wb.save(path)
                with self.assertRaises(ValueError):
                    audio.collect_items(path)

    def test_unknown_pronunciation_override_fails(self):
        with self.assertRaises(ValueError):
            audio.make_plan([], "source", {"w999999": "noun"})


if __name__ == "__main__":
    unittest.main()
